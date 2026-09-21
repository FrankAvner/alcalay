# -*- coding: utf-8 -*-

"""
Alcalay - Gmail INDEX

Indexes already copied and parsed Gmail messages.

Architecture:

    Gmail
      |
      v
    COPY
      |
      v
    storage/gmail/<account>/messages/<message_key>/message.eml
      |
      v
    PARSE
      |
      +----> gmail_parsed_messages
      |
      +----> gmail_attachments
      |
      v
    INDEX
      |
      +----> gmail_search_index
      |
      +----> gmail_attachment_search_index

Important:

- This module does NOT connect to Gmail.
- This module does NOT download messages from Gmail.
- This module does NOT modify message.eml files.
- Email content is indexed from gmail_parsed_messages.
- Attachment content is indexed from gmail_attachments/local files.
- Existing indexed records are skipped unless --force is used.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import sys

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_ROOT = PROJECT_ROOT / "src"

if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

DEFAULT_STORAGE_ROOT = PROJECT_ROOT / "storage" / "gmail"
DEFAULT_ACCOUNT = "frank.avner@gmail.com"

INDEX_VERSION = "2.1"
ATTACHMENT_INDEX_VERSION = "1.1"


class GmailIndexer:
    """
    Gmail email and attachment indexer.

    The indexer works only with local data and PostgreSQL.
    """

    def __init__(
        self,
        account_email: str,
        storage_root: Optional[Path] = None,
        force: bool = False,
        limit: Optional[int] = None,
    ):
        self.account_email = account_email

        self.storage_root = (
            Path(storage_root)
            if storage_root is not None
            else DEFAULT_STORAGE_ROOT
        )

        self.force = force
        self.limit = limit

        self.db = None
        self.gmail_account_id = None

        self.stats = {
            "messages_found": 0,
            "messages_indexed": 0,
            "messages_already_indexed": 0,
            "messages_updated": 0,
            "messages_missing": 0,
            "messages_failed": 0,
            "attachments_found": 0,
            "attachments_indexed": 0,
            "attachments_already_indexed": 0,
            "attachments_updated": 0,
            "attachments_with_text": 0,
            "attachments_without_text": 0,
            "attachments_unsupported": 0,
            "attachments_missing": 0,
            "attachments_failed": 0,
        }

    # ------------------------------------------------------------------
    # General helpers
    # ------------------------------------------------------------------

    def log(self, message: str) -> None:
        print(message, flush=True)

    def safe_json(self, value: Any) -> Any:
        if value is None:
            return {}

        if isinstance(value, (dict, list)):
            return value

        if isinstance(value, str):
            try:
                return json.loads(value)
            except Exception:
                return {"value": value}

        return {"value": str(value)}

    def json_value(self, value: Any) -> str:
        return json.dumps(
            value if value is not None else {},
            ensure_ascii=False,
        )

    def now_utc(self) -> datetime:
        return datetime.now(timezone.utc)

    def clean_text(self, value: Any) -> str:
        if value is None:
            return ""

        text = str(value)

        text = text.replace("\x00", " ")
        text = html.unescape(text)

        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n[ \t]+", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)

        return text.strip()

    def calculate_sha256(self, data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    # ------------------------------------------------------------------
    # Database
    # ------------------------------------------------------------------

    def load_database_context(self) -> None:
        from database.connection import DatabaseConnection

        connection = DatabaseConnection()
        self.db = connection.connect()

        with self.db.cursor() as cursor:
            cursor.execute(
                """
                SELECT id
                FROM gmail_accounts
                WHERE email = %s
                LIMIT 1
                """,
                (self.account_email,),
            )

            row = cursor.fetchone()

        if not row:
            raise RuntimeError(
                f"Gmail account not found in PostgreSQL: "
                f"{self.account_email}"
            )

        self.gmail_account_id = row[0]

        self.log(
            f"[DB] Gmail account ID: {self.gmail_account_id}"
        )

    def ensure_email_index_table(self) -> None:
        self.log("[DB] Ensuring email index table...")

        with self.db.cursor() as cursor:
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS gmail_search_index (
                    id BIGSERIAL PRIMARY KEY,
                    gmail_message_id BIGINT NOT NULL,
                    gmail_message_key TEXT NOT NULL,
                    gmail_account_id BIGINT NOT NULL,
                    mime_message_id TEXT,
                    thread_id TEXT,
                    subject TEXT,
                    sender_name TEXT,
                    sender_email TEXT,
                    recipients JSONB,
                    cc_recipients JSONB,
                    bcc_recipients JSONB,
                    reply_to JSONB,
                    date_sent TIMESTAMPTZ,
                    body_text TEXT,
                    body_html TEXT,
                    search_text TEXT,
                    has_attachments BOOLEAN DEFAULT FALSE,
                    attachment_count INTEGER DEFAULT 0,
                    raw_sha256 TEXT,
                    raw_size BIGINT,
                    parser_version TEXT,
                    index_version TEXT NOT NULL,
                    indexed_at TIMESTAMPTZ NOT NULL,
                    updated_at TIMESTAMPTZ NOT NULL,
                    search_vector TSVECTOR,
                    metadata JSONB
                )
                """
            )

            cursor.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS
                ux_gmail_search_index_message
                ON gmail_search_index (gmail_message_id)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_search_index_account
                ON gmail_search_index (gmail_account_id)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_search_index_date
                ON gmail_search_index (date_sent)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_search_index_message_key
                ON gmail_search_index (gmail_message_key)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_search_index_sender
                ON gmail_search_index (sender_email)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_search_index_subject
                ON gmail_search_index (subject)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_search_index_search_vector
                ON gmail_search_index
                USING GIN (search_vector)
                """
            )

        self.db.commit()

    def ensure_attachment_index_table(self) -> None:
        self.log("[DB] Ensuring attachment index table...")

        with self.db.cursor() as cursor:
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS gmail_attachment_search_index (
                    id BIGSERIAL PRIMARY KEY,
                    gmail_attachment_id BIGINT NOT NULL,
                    gmail_message_id BIGINT NOT NULL,
                    gmail_message_key TEXT,
                    gmail_account_id BIGINT NOT NULL,
                    message_id TEXT,
                    thread_id TEXT,
                    file_name TEXT,
                    mime_type TEXT,
                    file_size BIGINT,
                    content_hash TEXT,
                    local_file_id TEXT,
                    local_path TEXT,
                    extracted_text TEXT,
                    search_text TEXT,
                    extraction_method TEXT,
                    extraction_status TEXT,
                    extraction_error TEXT,
                    index_version TEXT NOT NULL,
                    indexed_at TIMESTAMPTZ NOT NULL,
                    updated_at TIMESTAMPTZ NOT NULL,
                    search_vector TSVECTOR,
                    metadata JSONB
                )
                """
            )

            cursor.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS
                ux_gmail_attachment_search_index_attachment
                ON gmail_attachment_search_index (gmail_attachment_id)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_attachment_search_index_account
                ON gmail_attachment_search_index (gmail_account_id)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_attachment_search_index_message
                ON gmail_attachment_search_index (gmail_message_id)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_attachment_search_index_hash
                ON gmail_attachment_search_index (content_hash)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_attachment_search_index_filename
                ON gmail_attachment_search_index (file_name)
                """
            )

            cursor.execute(
                """
                CREATE INDEX IF NOT EXISTS
                ix_gmail_attachment_search_index_search_vector
                ON gmail_attachment_search_index
                USING GIN (search_vector)
                """
            )

        self.db.commit()

    # ------------------------------------------------------------------
    # Email loading
    # ------------------------------------------------------------------

    def get_parsed_messages(self) -> list[dict[str, Any]]:
        self.log("[INDEX] Loading parsed messages...")

        sql = """
            SELECT
                gpm.id,
                gpm.gmail_message_id,
                gpm.gmail_message_key,
                gpm.mime_message_id,
                gpm.thread_id,
                gpm.subject,
                gpm.sender_name,
                gpm.sender_email,
                gpm.recipients,
                gpm.cc_recipients,
                gpm.bcc_recipients,
                gpm.reply_to,
                gpm.date_header,
                gpm.date_sent,
                gpm.in_reply_to,
                gpm.references_header,
                gpm.content_type,
                gpm.charset,
                gpm.body_text,
                gpm.body_html,
                gpm.search_text,
                gpm.raw_sha256,
                gpm.raw_size,
                gpm.has_attachments,
                gpm.attachment_count,
                gpm.headers,
                gpm.metadata,
                gpm.parser_version,
                gpm.parsed_at,
                gm.message_id,
                gm.history_id,
                gm.sender AS gmail_sender,
                gm.recipients AS gmail_recipients,
                gm.cc AS gmail_cc,
                gm.bcc AS gmail_bcc,
                gm.received_at
            FROM gmail_parsed_messages gpm
            INNER JOIN gmail_messages gm
                ON gm.id = gpm.gmail_message_id
            WHERE gm.gmail_account_id = %s
            ORDER BY gpm.id
        """

        parameters: list[Any] = [
            self.gmail_account_id
        ]

        if self.limit is not None:
            sql += " LIMIT %s"
            parameters.append(self.limit)

        with self.db.cursor() as cursor:
            cursor.execute(sql, parameters)

            rows = cursor.fetchall()

            columns = [
                description.name
                for description in cursor.description
            ]

        messages = [
            dict(zip(columns, row))
            for row in rows
        ]

        self.stats["messages_found"] = len(messages)

        self.log(
            f"[INDEX] Parsed messages found: "
            f"{len(messages)}"
        )

        return messages

    # ------------------------------------------------------------------
    # Email index checks
    # ------------------------------------------------------------------

    def get_existing_email_index(
        self,
        gmail_message_id: int,
    ) -> Optional[dict[str, Any]]:
        with self.db.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    id,
                    index_version,
                    raw_sha256,
                    parser_version
                FROM gmail_search_index
                WHERE gmail_message_id = %s
                LIMIT 1
                """,
                (gmail_message_id,),
            )

            row = cursor.fetchone()

        if not row:
            return None

        return {
            "id": row[0],
            "index_version": row[1],
            "raw_sha256": row[2],
            "parser_version": row[3],
        }

    # ------------------------------------------------------------------
    # Email indexing
    # ------------------------------------------------------------------

    def build_email_search_text(
        self,
        message: dict[str, Any],
    ) -> str:
        parts: list[str] = []

        fields = [
            message.get("subject"),
            message.get("sender_name"),
            message.get("sender_email"),
            message.get("recipients"),
            message.get("cc_recipients"),
            message.get("bcc_recipients"),
            message.get("reply_to"),
            message.get("body_text"),
            message.get("body_html"),
            message.get("search_text"),
        ]

        for value in fields:
            if value is None:
                continue

            if isinstance(value, (dict, list)):
                value = json.dumps(
                    value,
                    ensure_ascii=False,
                )

            cleaned = self.clean_text(value)

            if cleaned:
                parts.append(cleaned)

        return "\n".join(parts)

    def index_email(
        self,
        message: dict[str, Any],
    ) -> str:
        gmail_message_id = message["gmail_message_id"]

        existing = self.get_existing_email_index(
            gmail_message_id
        )

        raw_sha256 = message.get("raw_sha256")

        if (
            existing
            and not self.force
            and existing["index_version"] == INDEX_VERSION
            and existing["raw_sha256"] == raw_sha256
        ):
            return "already_indexed"

        search_text = self.build_email_search_text(
            message
        )

        metadata = self.safe_json(
            message.get("metadata")
        )

        if not isinstance(metadata, dict):
            metadata = {
                "value": metadata
            }

        metadata["index_version"] = INDEX_VERSION
        metadata["indexed_by"] = "gmail_indexer"

        indexed_at = self.now_utc()

        recipients = self.safe_json(
            message.get("recipients")
        )

        cc_recipients = self.safe_json(
            message.get("cc_recipients")
        )

        bcc_recipients = self.safe_json(
            message.get("bcc_recipients")
        )

        reply_to = self.safe_json(
            message.get("reply_to")
        )

        has_attachments = bool(
            message.get("has_attachments")
        )

        attachment_count = int(
            message.get("attachment_count") or 0
        )

        parser_version = message.get(
            "parser_version"
        )

        body_text = message.get("body_text") or ""
        body_html = message.get("body_html") or ""
        subject = message.get("subject") or ""

        sender_name = (
            message.get("sender_name") or ""
        )

        sender_email = (
            message.get("sender_email") or ""
        )

        mime_message_id = message.get(
            "mime_message_id"
        )

        thread_id = message.get(
            "thread_id"
        )

        gmail_message_key = (
            message.get("gmail_message_key")
            or message.get("message_id")
            or ""
        )

        date_sent = message.get(
            "date_sent"
        )

        raw_size = message.get(
            "raw_size"
        )

        parameters = (
            gmail_message_id,
            gmail_message_key,
            self.gmail_account_id,
            mime_message_id,
            thread_id,
            subject,
            sender_name,
            sender_email,
            self.json_value(recipients),
            self.json_value(cc_recipients),
            self.json_value(bcc_recipients),
            self.json_value(reply_to),
            date_sent,
            body_text,
            body_html,
            search_text,
            has_attachments,
            attachment_count,
            raw_sha256,
            raw_size,
            parser_version,
            INDEX_VERSION,
            indexed_at,
            indexed_at,
            search_text,
            self.json_value(metadata),
        )

        with self.db.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO gmail_search_index (
                    gmail_message_id,
                    gmail_message_key,
                    gmail_account_id,
                    mime_message_id,
                    thread_id,
                    subject,
                    sender_name,
                    sender_email,
                    recipients,
                    cc_recipients,
                    bcc_recipients,
                    reply_to,
                    date_sent,
                    body_text,
                    body_html,
                    search_text,
                    has_attachments,
                    attachment_count,
                    raw_sha256,
                    raw_size,
                    parser_version,
                    index_version,
                    indexed_at,
                    updated_at,
                    search_vector,
                    metadata
                )
                VALUES (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s::jsonb,
                    %s::jsonb,
                    %s::jsonb,
                    %s::jsonb,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    to_tsvector(
                        'simple',
                        coalesce(%s, '')
                    ),
                    %s::jsonb
                )
                ON CONFLICT (gmail_message_id)
                DO UPDATE SET
                    gmail_message_key = EXCLUDED.gmail_message_key,
                    gmail_account_id = EXCLUDED.gmail_account_id,
                    mime_message_id = EXCLUDED.mime_message_id,
                    thread_id = EXCLUDED.thread_id,
                    subject = EXCLUDED.subject,
                    sender_name = EXCLUDED.sender_name,
                    sender_email = EXCLUDED.sender_email,
                    recipients = EXCLUDED.recipients,
                    cc_recipients = EXCLUDED.cc_recipients,
                    bcc_recipients = EXCLUDED.bcc_recipients,
                    reply_to = EXCLUDED.reply_to,
                    date_sent = EXCLUDED.date_sent,
                    body_text = EXCLUDED.body_text,
                    body_html = EXCLUDED.body_html,
                    search_text = EXCLUDED.search_text,
                    has_attachments = EXCLUDED.has_attachments,
                    attachment_count = EXCLUDED.attachment_count,
                    raw_sha256 = EXCLUDED.raw_sha256,
                    raw_size = EXCLUDED.raw_size,
                    parser_version = EXCLUDED.parser_version,
                    index_version = EXCLUDED.index_version,
                    indexed_at = EXCLUDED.indexed_at,
                    updated_at = EXCLUDED.updated_at,
                    search_vector = EXCLUDED.search_vector,
                    metadata = EXCLUDED.metadata
                """,
                parameters,
            )

        self.db.commit()

        if existing:
            return "updated"

        return "indexed"

    def process_email(
        self,
        position: int,
        total: int,
        message: dict[str, Any],
    ) -> None:
        message_key = (
            message.get("gmail_message_key")
            or message.get("message_id")
            or ""
        )

        subject = (
            message.get("subject")
            or ""
        ).strip()

        self.log(
            f"[MESSAGE] {position}/{total} | "
            f"{message_key} | {subject}"
        )

        try:
            result = self.index_email(message)

            if result == "indexed":
                self.stats["messages_indexed"] += 1
                self.log("[MESSAGE] indexed")

            elif result == "already_indexed":
                self.stats[
                    "messages_already_indexed"
                ] += 1
                self.log(
                    "[MESSAGE] already_indexed"
                )

            elif result == "updated":
                self.stats["messages_updated"] += 1
                self.log("[MESSAGE] updated")

            else:
                self.stats["messages_failed"] += 1
                self.log(
                    f"[MESSAGE] ERROR: "
                    f"unknown result: {result}"
                )

        except Exception as exc:
            self.stats["messages_failed"] += 1

            self.log(
                f"[MESSAGE] ERROR: {exc}"
            )

            self.db.rollback()

    # ------------------------------------------------------------------
    # Attachment loading
    # ------------------------------------------------------------------

    def get_attachments(self) -> list[dict[str, Any]]:
        self.log(
            "[INDEX] Loading existing attachments..."
        )

        sql = """
            SELECT
                ga.id,
                ga.gmail_message_id,
                ga.attachment_id,
                ga.file_name,
                ga.mime_type,
                ga.file_size,
                ga.content_hash,
                ga.local_file_id,
                ga.saved_to_local,
                ga.metadata,
                gm.gmail_account_id,
                gm.message_id,
                gm.thread_id
            FROM gmail_attachments ga
            INNER JOIN gmail_messages gm
                ON gm.id = ga.gmail_message_id
            WHERE gm.gmail_account_id = %s
            ORDER BY ga.id
        """

        with self.db.cursor() as cursor:
            cursor.execute(
                sql,
                (self.gmail_account_id,),
            )

            rows = cursor.fetchall()

            columns = [
                description.name
                for description in cursor.description
            ]

        attachments = [
            dict(zip(columns, row))
            for row in rows
        ]

        self.stats["attachments_found"] = len(
            attachments
        )

        self.log(
            "[INDEX] Existing attachments found: "
            f"{len(attachments)}"
        )

        return attachments

    # ------------------------------------------------------------------
    # Attachment paths
    # ------------------------------------------------------------------

    def get_attachment_local_path(
        self,
        attachment: dict[str, Any],
    ) -> Optional[Path]:
        metadata = self.safe_json(
            attachment.get("metadata")
        )

        if isinstance(metadata, dict):
            local_path = metadata.get(
                "local_path"
            )

            if local_path:
                path = Path(str(local_path))

                if path.exists() and path.is_file():
                    return path

                if not path.is_absolute():
                    candidate = (
                        PROJECT_ROOT / path
                    )

                    if (
                        candidate.exists()
                        and candidate.is_file()
                    ):
                        return candidate

        local_file_id = attachment.get(
            "local_file_id"
        )

        if local_file_id:
            path = Path(
                str(local_file_id)
            )

            if path.exists() and path.is_file():
                return path

            if not path.is_absolute():
                candidate = (
                    PROJECT_ROOT / path
                )

                if (
                    candidate.exists()
                    and candidate.is_file()
                ):
                    return candidate

        return None

    # ------------------------------------------------------------------
    # Attachment extraction
    # ------------------------------------------------------------------

    def extract_text_from_txt(
        self,
        path: Path,
    ) -> tuple[str, str]:
        data = path.read_bytes()

        encodings = [
            "utf-8",
            "utf-8-sig",
            "cp1255",
            "cp1252",
            "latin-1",
        ]

        for encoding in encodings:
            try:
                text = data.decode(encoding)

                return (
                    self.clean_text(text),
                    f"text:{encoding}",
                )

            except UnicodeDecodeError:
                continue

        return "", "text:failed"

    def extract_text_from_csv(
        self,
        path: Path,
    ) -> tuple[str, str]:
        data = path.read_bytes()

        encodings = [
            "utf-8",
            "utf-8-sig",
            "cp1255",
            "cp1252",
            "latin-1",
        ]

        text = None

        for encoding in encodings:
            try:
                text = data.decode(encoding)
                break

            except UnicodeDecodeError:
                continue

        if text is None:
            return "", "csv:failed"

        lines: list[str] = []

        try:
            reader = csv.reader(
                text.splitlines()
            )

            for row in reader:
                lines.append(
                    " ".join(
                        str(value)
                        for value in row
                    )
                )

            return (
                self.clean_text(
                    "\n".join(lines)
                ),
                "csv",
            )

        except Exception:
            return (
                self.clean_text(text),
                "csv:raw",
            )

    def extract_text_from_html(
        self,
        path: Path,
    ) -> tuple[str, str]:
        try:
            data = path.read_bytes()

            text = data.decode(
                "utf-8",
                errors="replace",
            )

            text = re.sub(
                r"<script\b[^>]*>.*?</script>",
                " ",
                text,
                flags=re.IGNORECASE | re.DOTALL,
            )

            text = re.sub(
                r"<style\b[^>]*>.*?</style>",
                " ",
                text,
                flags=re.IGNORECASE | re.DOTALL,
            )

            text = re.sub(
                r"<[^>]+>",
                " ",
                text,
            )

            return (
                self.clean_text(text),
                "html",
            )

        except Exception:
            return "", "html:failed"

    def extract_text_from_pdf(
        self,
        path: Path,
    ) -> tuple[str, str]:
        try:
            from pypdf import PdfReader

        except ImportError:
            return (
                "",
                "pdf:pypdf_not_installed",
            )

        try:
            reader = PdfReader(
                str(path)
            )

            pages: list[str] = []

            for page in reader.pages:
                try:
                    text = page.extract_text()
                except Exception:
                    text = ""

                if text:
                    pages.append(text)

            return (
                self.clean_text(
                    "\n\n".join(pages)
                ),
                "pdf:pypdf",
            )

        except Exception as exc:
            return (
                "",
                f"pdf:error:{exc}",
            )

    def extract_text_from_docx(
        self,
        path: Path,
    ) -> tuple[str, str]:
        try:
            from docx import Document

        except ImportError:
            return (
                "",
                "docx:python-docx_not_installed",
            )

        try:
            document = Document(
                str(path)
            )

            parts: list[str] = []

            for paragraph in document.paragraphs:
                if paragraph.text:
                    parts.append(
                        paragraph.text
                    )

            for table in document.tables:
                for row in table.rows:
                    values = []

                    for cell in row.cells:
                        values.append(
                            cell.text
                        )

                    parts.append(
                        " | ".join(values)
                    )

            return (
                self.clean_text(
                    "\n".join(parts)
                ),
                "docx:python-docx",
            )

        except Exception as exc:
            return (
                "",
                f"docx:error:{exc}",
            )

    def extract_text_from_xlsx(
        self,
        path: Path,
    ) -> tuple[str, str]:
        try:
            from openpyxl import load_workbook

        except ImportError:
            return (
                "",
                "xlsx:openpyxl_not_installed",
            )

        try:
            workbook = load_workbook(
                filename=str(path),
                read_only=True,
                data_only=True,
            )

            parts: list[str] = []

            for worksheet in workbook.worksheets:
                parts.append(
                    f"[SHEET] {worksheet.title}"
                )

                for row in worksheet.iter_rows(
                    values_only=True
                ):
                    values = []

                    for value in row:
                        if value is None:
                            values.append("")
                        else:
                            values.append(
                                str(value)
                            )

                    line = " | ".join(values)

                    if line.strip():
                        parts.append(line)

            workbook.close()

            return (
                self.clean_text(
                    "\n".join(parts)
                ),
                "xlsx:openpyxl",
            )

        except Exception as exc:
            return (
                "",
                f"xlsx:error:{exc}",
            )

    def extract_text_from_pptx(
        self,
        path: Path,
    ) -> tuple[str, str]:
        try:
            from pptx import Presentation

        except ImportError:
            return (
                "",
                "pptx:python-pptx_not_installed",
            )

        try:
            presentation = Presentation(
                str(path)
            )

            parts: list[str] = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1,
            ):
                parts.append(
                    f"[SLIDE {slide_number}]"
                )

                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text = shape.text

                        if text:
                            parts.append(text)

            return (
                self.clean_text(
                    "\n".join(parts)
                ),
                "pptx:python-pptx",
            )

        except Exception as exc:
            return (
                "",
                f"pptx:error:{exc}",
            )

    def extract_text_from_rtf(
        self,
        path: Path,
    ) -> tuple[str, str]:
        try:
            raw = path.read_text(
                encoding="utf-8",
                errors="ignore",
            )

            raw = re.sub(
                r"\\'[0-9a-fA-F]{2}",
                " ",
                raw,
            )

            raw = re.sub(
                r"\\[a-zA-Z]+-?\d*\*? ?",
                " ",
                raw,
            )

            raw = raw.replace(
                "{",
                " ",
            ).replace(
                "}",
                " ",
            )

            return (
                self.clean_text(raw),
                "rtf:basic",
            )

        except Exception:
            return "", "rtf:failed"

    def extract_attachment_text(
        self,
        path: Path,
        file_name: str,
        mime_type: str,
    ) -> tuple[str, str, str]:
        suffix = path.suffix.lower()
        mime_lower = mime_type.lower()

        if (
            suffix == ".pdf"
            or mime_lower == "application/pdf"
        ):
            text, method = (
                self.extract_text_from_pdf(path)
            )

            return text, method, "pdf"

        if suffix == ".docx":
            text, method = (
                self.extract_text_from_docx(path)
            )

            return text, method, "docx"

        if suffix == ".xlsx":
            text, method = (
                self.extract_text_from_xlsx(path)
            )

            return text, method, "xlsx"

        if suffix == ".pptx":
            text, method = (
                self.extract_text_from_pptx(path)
            )

            return text, method, "pptx"

        if suffix in {
            ".txt",
            ".text",
            ".log",
            ".md",
            ".json",
            ".xml",
            ".yaml",
            ".yml",
        }:
            text, method = (
                self.extract_text_from_txt(path)
            )

            return text, method, "text"

        if suffix in {
            ".csv",
            ".tsv",
        }:
            text, method = (
                self.extract_text_from_csv(path)
            )

            return text, method, "csv"

        if suffix in {
            ".html",
            ".htm",
        }:
            text, method = (
                self.extract_text_from_html(path)
            )

            return text, method, "html"

        if suffix == ".rtf":
            text, method = (
                self.extract_text_from_rtf(path)
            )

            return text, method, "rtf"

        if (
            mime_lower.startswith("text/")
            or "html" in mime_lower
        ):
            text, method = (
                self.extract_text_from_txt(path)
            )

            return text, method, "text"

        return "", "unsupported", "unsupported"

    # ------------------------------------------------------------------
    # Attachment index checks
    # ------------------------------------------------------------------

    def get_existing_attachment_index(
        self,
        attachment_id: int,
    ) -> Optional[dict[str, Any]]:
        with self.db.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    id,
                    index_version,
                    content_hash
                FROM gmail_attachment_search_index
                WHERE gmail_attachment_id = %s
                LIMIT 1
                """,
                (attachment_id,),
            )

            row = cursor.fetchone()

        if not row:
            return None

        return {
            "id": row[0],
            "index_version": row[1],
            "content_hash": row[2],
        }

    # ------------------------------------------------------------------
    # Attachment indexing
    # ------------------------------------------------------------------

    def index_attachment(
        self,
        attachment: dict[str, Any],
    ) -> str:
        attachment_id = attachment["id"]

        existing = (
            self.get_existing_attachment_index(
                attachment_id
            )
        )

        content_hash = attachment.get(
            "content_hash"
        )

        if (
            existing
            and not self.force
            and existing["index_version"]
            == ATTACHMENT_INDEX_VERSION
            and existing["content_hash"]
            == content_hash
        ):
            return "already_indexed"

        file_name = (
            attachment.get("file_name")
            or ""
        )

        mime_type = (
            attachment.get("mime_type")
            or ""
        )

        local_path = (
            self.get_attachment_local_path(
                attachment
            )
        )

        if local_path is None:
            self.stats[
                "attachments_missing"
            ] += 1

            extraction_status = "missing_file"
            extraction_method = ""
            extraction_error = (
                "Local attachment file not found"
            )
            extracted_text = ""

        else:
            try:
                (
                    extracted_text,
                    extraction_method,
                    extraction_type,
                ) = self.extract_attachment_text(
                    local_path,
                    file_name,
                    mime_type,
                )

                if extraction_type == "unsupported":
                    self.stats[
                        "attachments_unsupported"
                    ] += 1

                    extraction_status = "unsupported"
                    extraction_error = None

                elif extracted_text.strip():
                    self.stats[
                        "attachments_with_text"
                    ] += 1

                    extraction_status = (
                        "text_extracted"
                    )
                    extraction_error = None

                else:
                    self.stats[
                        "attachments_without_text"
                    ] += 1

                    extraction_status = "no_text"
                    extraction_error = None

            except Exception as exc:
                self.stats[
                    "attachments_failed"
                ] += 1

                extraction_status = "failed"
                extraction_method = ""
                extraction_error = str(exc)
                extracted_text = ""

        metadata = self.safe_json(
            attachment.get("metadata")
        )

        if not isinstance(metadata, dict):
            metadata = {
                "value": metadata
            }

        metadata["index_version"] = (
            ATTACHMENT_INDEX_VERSION
        )

        metadata["indexed_by"] = (
            "gmail_indexer"
        )

        if local_path is not None:
            metadata["indexed_local_path"] = (
                str(local_path)
            )

        search_parts = [
            file_name,
            mime_type,
            extracted_text,
        ]

        search_text = self.clean_text(
            "\n".join(
                part
                for part in search_parts
                if part
            )
        )

        indexed_at = self.now_utc()

        parameters = (
            attachment_id,
            attachment.get(
                "gmail_message_id"
            ),
            attachment.get(
                "message_id"
            ) or "",
            self.gmail_account_id,
            attachment.get(
                "message_id"
            ),
            attachment.get(
                "thread_id"
            ),
            file_name,
            mime_type,
            attachment.get(
                "file_size"
            ),
            content_hash,
            attachment.get(
                "local_file_id"
            ),
            (
                str(local_path)
                if local_path is not None
                else None
            ),
            extracted_text,
            search_text,
            extraction_method,
            extraction_status,
            extraction_error,
            ATTACHMENT_INDEX_VERSION,
            indexed_at,
            indexed_at,
            search_text,
            self.json_value(metadata),
        )

        with self.db.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO gmail_attachment_search_index (
                    gmail_attachment_id,
                    gmail_message_id,
                    gmail_message_key,
                    gmail_account_id,
                    message_id,
                    thread_id,
                    file_name,
                    mime_type,
                    file_size,
                    content_hash,
                    local_file_id,
                    local_path,
                    extracted_text,
                    search_text,
                    extraction_method,
                    extraction_status,
                    extraction_error,
                    index_version,
                    indexed_at,
                    updated_at,
                    search_vector,
                    metadata
                )
                VALUES (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    to_tsvector(
                        'simple',
                        coalesce(%s, '')
                    ),
                    %s::jsonb
                )
                ON CONFLICT (gmail_attachment_id)
                DO UPDATE SET
                    gmail_message_id = EXCLUDED.gmail_message_id,
                    gmail_message_key = EXCLUDED.gmail_message_key,
                    gmail_account_id = EXCLUDED.gmail_account_id,
                    message_id = EXCLUDED.message_id,
                    thread_id = EXCLUDED.thread_id,
                    file_name = EXCLUDED.file_name,
                    mime_type = EXCLUDED.mime_type,
                    file_size = EXCLUDED.file_size,
                    content_hash = EXCLUDED.content_hash,
                    local_file_id = EXCLUDED.local_file_id,
                    local_path = EXCLUDED.local_path,
                    extracted_text = EXCLUDED.extracted_text,
                    search_text = EXCLUDED.search_text,
                    extraction_method = EXCLUDED.extraction_method,
                    extraction_status = EXCLUDED.extraction_status,
                    extraction_error = EXCLUDED.extraction_error,
                    index_version = EXCLUDED.index_version,
                    indexed_at = EXCLUDED.indexed_at,
                    updated_at = EXCLUDED.updated_at,
                    search_vector = EXCLUDED.search_vector,
                    metadata = EXCLUDED.metadata
                """,
                parameters,
            )

        self.db.commit()

        if existing:
            return "updated"

        return "indexed"

    def process_attachment(
        self,
        position: int,
        total: int,
        attachment: dict[str, Any],
    ) -> None:
        file_name = (
            attachment.get("file_name")
            or ""
        )

        self.log(
            f"[ATTACHMENT] "
            f"{position}/{total} | {file_name}"
        )

        try:
            result = self.index_attachment(
                attachment
            )

            if result == "indexed":
                self.stats[
                    "attachments_indexed"
                ] += 1

                self.log(
                    f"[ATTACHMENT] indexed | "
                    f"{file_name}"
                )

            elif result == "already_indexed":
                self.stats[
                    "attachments_already_indexed"
                ] += 1

                self.log(
                    f"[ATTACHMENT] already_indexed | "
                    f"{file_name}"
                )

            elif result == "updated":
                self.stats[
                    "attachments_updated"
                ] += 1

                self.log(
                    f"[ATTACHMENT] updated | "
                    f"{file_name}"
                )

            else:
                self.log(
                    f"[ATTACHMENT] result={result} | "
                    f"{file_name}"
                )

        except Exception as exc:
            self.stats[
                "attachments_failed"
            ] += 1

            self.db.rollback()

            self.log(
                f"[ATTACHMENT] ERROR | "
                f"{file_name} | {exc}"
            )

    # ------------------------------------------------------------------
    # Run
    # ------------------------------------------------------------------

    def run(self) -> int:
        print()
        print("=" * 72)
        print("ALCALAY - GMAIL INDEX")
        print()
        print(
            f"Account: {self.account_email}"
        )
        print(
            f"Mode: "
            f"{'FORCE' if self.force else 'NORMAL'}"
        )
        print(
            f"Storage: {self.storage_root}"
        )
        print("=" * 72)
        print()

        try:
            self.load_database_context()

            self.ensure_email_index_table()
            self.ensure_attachment_index_table()

            messages = self.get_parsed_messages()

            total_messages = len(messages)

            for position, message in enumerate(
                messages,
                start=1,
            ):
                self.process_email(
                    position,
                    total_messages,
                    message,
                )

            attachments = self.get_attachments()

            total_attachments = len(
                attachments
            )

            for position, attachment in enumerate(
                attachments,
                start=1,
            ):
                self.process_attachment(
                    position,
                    total_attachments,
                    attachment,
                )

            self.print_summary()

            return 0

        except KeyboardInterrupt:
            print()
            print(
                "[INDEX] Interrupted by user."
            )
            return 130

        except Exception as exc:
            print()
            print(
                f"[INDEX] FATAL ERROR: {exc}"
            )

            if self.db is not None:
                try:
                    self.db.rollback()
                except Exception:
                    pass

            return 1

        finally:
            if self.db is not None:
                try:
                    self.db.close()
                except Exception:
                    pass

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------

    def print_summary(self) -> None:
        print()
        print("=" * 72)
        print("ALCALAY - GMAIL INDEX SUMMARY")
        print("=" * 72)
        print()

        print(
            f"Account: {self.account_email}"
        )

        print()

        print("EMAILS")

        print(
            f"  Found: "
            f"{self.stats['messages_found']}"
        )

        print(
            f"  Indexed new: "
            f"{self.stats['messages_indexed']}"
        )

        print(
            f"  Already indexed: "
            f"{self.stats['messages_already_indexed']}"
        )

        print(
            f"  Updated: "
            f"{self.stats['messages_updated']}"
        )

        print(
            f"  Missing: "
            f"{self.stats['messages_missing']}"
        )

        print(
            f"  Failed: "
            f"{self.stats['messages_failed']}"
        )

        print()

        print("ATTACHMENTS")

        print(
            f"  Found: "
            f"{self.stats['attachments_found']}"
        )

        print(
            f"  Indexed new: "
            f"{self.stats['attachments_indexed']}"
        )

        print(
            f"  Already indexed: "
            f"{self.stats['attachments_already_indexed']}"
        )

        print(
            f"  Updated: "
            f"{self.stats['attachments_updated']}"
        )

        print(
            f"  With extracted text: "
            f"{self.stats['attachments_with_text']}"
        )

        print(
            f"  Without extracted text: "
            f"{self.stats['attachments_without_text']}"
        )

        print(
            f"  Unsupported: "
            f"{self.stats['attachments_unsupported']}"
        )

        print(
            f"  Missing local file: "
            f"{self.stats['attachments_missing']}"
        )

        print(
            f"  Failed: "
            f"{self.stats['attachments_failed']}"
        )

        print("=" * 72)
        print()


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Alcalay Gmail email and attachment indexer"
        )
    )

    parser.add_argument(
        "--account",
        default=DEFAULT_ACCOUNT,
        help=(
            "Gmail account email "
            f"(default: {DEFAULT_ACCOUNT})"
        ),
    )

    parser.add_argument(
        "--storage-root",
        default=str(DEFAULT_STORAGE_ROOT),
        help=(
            "Gmail storage root "
            f"(default: {DEFAULT_STORAGE_ROOT})"
        ),
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help=(
            "Limit number of parsed messages "
            "processed in this run"
        ),
    )

    parser.add_argument(
        "--force",
        action="store_true",
        help=(
            "Re-index existing email and "
            "attachment records"
        ),
    )

    return parser.parse_args()


def main() -> int:
    args = parse_arguments()

    if args.limit is not None and args.limit <= 0:
        print(
            "ERROR: --limit must be greater than zero."
        )
        return 2

    indexer = GmailIndexer(
        account_email=args.account,
        storage_root=Path(
            args.storage_root
        ),
        force=args.force,
        limit=args.limit,
    )

    return indexer.run()


if __name__ == "__main__":
    sys.exit(main())